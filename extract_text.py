import sys
import os

try:
    from docx import Document
    import pptx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "python-pptx", "fpdf2"])
    from docx import Document
    import pptx

def extract_docx(file_path):
    doc = Document(file_path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            text.append("\t".join(row_data))
    return "\n".join(text)

def extract_pptx(file_path):
    prs = pptx.Presentation(file_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

with open("report_text.txt", "w", encoding="utf-8") as f:
    f.write(extract_docx("EasyDocs_Minor_Project_Report_FINAL.docx"))

with open("ppt_text.txt", "w", encoding="utf-8") as f:
    f.write(extract_pptx("EasyDocs_Formatted_Presentation.pptx"))
